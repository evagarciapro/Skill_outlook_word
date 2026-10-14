from docx import Document

def generate_word_report(emails, output_path):
    doc = Document()
    doc.add_heading("Compte rendu hebdomadaire", level=1)

    doc.add_paragraph(f"Nombre total d'e-mails analysés : {len(emails)}")

    doc.add_heading("Résumé des e-mails", level=2)

    for email in emails:
        doc.add_heading(email["subject"] or "(Sans objet)", level=3)
        doc.add_paragraph(f"Dossier : {email['folder']}")
        doc.add_paragraph(f"Date : {email['date']}")
        doc.add_paragraph(f"De : {email['sender']}")
        doc.add_paragraph(f"À : {email['to']}")
        doc.add_paragraph(email["body"][:1000])

    doc.save(output_path)